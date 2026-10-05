import sys
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter as L
sys.path.insert(0, '/home/claude/robots-txt-crawl')
from robots_eval import CRAWLERS

SP = '/tmp/claude-0/-home-claude/6c55c4bd-efa3-5ed9-a3f2-bd0d0dd5339b/scratchpad'
OUT = f'{SP}/robots_txt_top50k_crawler_access.xlsx'
d = pd.read_pickle(f'{SP}/final.pkl').sort_values('popularity_rank')
N = len(d)

F = 'Arial'
hdr_font = Font(name=F, bold=True, color='FFFFFF', size=10)
hdr_fill = PatternFill('solid', fgColor='1F3864')
co_fill = PatternFill('solid', fgColor='D9E1F2')
body = Font(name=F, size=10)
bold = Font(name=F, size=10, bold=True)
green = PatternFill('solid', fgColor='C6EFCE'); green_f = Font(name=F, size=10, color='006100')
red = PatternFill('solid', fgColor='FFC7CE'); red_f = Font(name=F, size=10, color='9C0006')
thin = Side(style='thin', color='BFBFBF')
ALLOW_FMT = '[=1]"Allowed";[=0]"Blocked";General'

crawl_cols = [c for c, _, _ in CRAWLERS] + ['All bots (*)']
company = {c: co for c, co, _ in CRAWLERS}; company['All bots (*)'] = 'Default rule'

info_cols = [
    ('Popularity rank (est.)', 'popularity_rank', 11), ('URL', 'origin', 34), ('Domain', 'registered_domain', 22),
    ('Page title', 'page_title', 34), ('Category', 'category', 22), ('Category source', 'category_source', 20),
    ('Country', 'country', 16), ('Country basis', 'country_method', 22), ('Top markets (est.)', 'top_markets', 13),
    ('Global reach', 'global_reach', 12), ('Chrome rank bucket (CrUX)', 'crux_bucket', 12), ('Tranco rank', 'tranco_rank', 10),
    ('robots.txt source', 'robots_source', 26),
]

wb = Workbook()
ws = wb.active; ws.title = 'Sites'
# Row 1: company grouping; Row 2: headers
for j, (h, _, w) in enumerate(info_cols, 1):
    ws.cell(2, j, h); ws.column_dimensions[L(j)].width = w
for k, c in enumerate(crawl_cols):
    j = len(info_cols) + 1 + k
    ws.cell(1, j, company[c]).font = Font(name=F, size=9, bold=True, color='1F3864')
    ws.cell(1, j).fill = co_fill; ws.cell(1, j).alignment = Alignment(horizontal='center')
    ws.cell(2, j, c); ws.column_dimensions[L(j)].width = 11
ws.cell(1, 1, 'Top 50,000 websites (Chrome UX Report, Aug 2026): robots.txt access by crawler. 1 = Allowed, 0 = Blocked, blank = no robots.txt data. See "Method & notes".').font = Font(name=F, size=9, italic=True)
for j in range(1, len(info_cols) + len(crawl_cols) + 1):
    c = ws.cell(2, j); c.font = hdr_font; c.fill = hdr_fill
    c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
ws.row_dimensions[2].height = 42

bucket_lbl = {1000: 'Top 1K', 5000: 'Top 5K', 10000: 'Top 10K', 50000: 'Top 50K'}
recs = d[[k for _, k, _ in info_cols] + crawl_cols].to_dict('records')
for i, rd in enumerate(recs, 3):
    for j, (_, key, _) in enumerate(info_cols, 1):
        v = rd[key]
        if key == 'crux_bucket': v = bucket_lbl[v]
        if key == 'tranco_rank': v = None if pd.isna(v) else int(v)
        ws.cell(i, j, v)
    for k, c in enumerate(crawl_cols):
        v = rd[c]
        ws.cell(i, len(info_cols) + 1 + k, None if pd.isna(v) else int(v))
last = N + 2
first_c = len(info_cols) + 1; last_c = len(info_cols) + len(crawl_cols)
rng = f'{L(first_c)}3:{L(last_c)}{last}'
for row in ws.iter_rows(min_row=3, max_row=last, min_col=1, max_col=last_c):
    for c in row:
        c.font = body
for row in ws.iter_rows(min_row=3, max_row=last, min_col=first_c, max_col=last_c):
    for c in row:
        c.number_format = ALLOW_FMT; c.alignment = Alignment(horizontal='center')
ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['1'], fill=green, font=green_f))
ws.conditional_formatting.add(rng, CellIsRule(operator='equal', formula=['0'], fill=red, font=red_f))
ws.freeze_panes = ws.cell(3, 3)
ws.auto_filter.ref = f'A2:{L(last_c)}{last}'
for j in (1, 11, 12):
    for row in ws.iter_rows(min_row=3, max_row=last, min_col=j, max_col=j):
        row[0].number_format = '#,##0'; row[0].alignment = Alignment(horizontal='center')

# ---------------- Summary ----------------
S = wb.create_sheet('Summary')
col_of = {c: L(first_c + k) for k, c in enumerate(crawl_cols)}
CAT = f"Sites!$E$3:$E${last}"; BUCKET = f"Sites!$K$3:$K${last}"
S['A1'] = 'Block rates by crawler'; S['A1'].font = Font(name=F, size=14, bold=True)
S['A2'] = 'Share of sites whose robots.txt blocks the crawler from the homepage / site content, among sites with robots.txt data. All cells are live formulas over the Sites tab.'
S['A2'].font = Font(name=F, size=9, italic=True)
heads = ['Company', 'Crawler (robots.txt token)', 'Sites with data', 'Sites blocking', '% blocked (all 50K)', '% blocked (Top 1K)', '% blocked (Top 10K)']
for j, h in enumerate(heads, 1):
    c = S.cell(4, j, h); c.font = hdr_font; c.fill = hdr_fill; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
S.row_dimensions[4].height = 32
for i, c in enumerate(crawl_cols, 5):
    col = f"Sites!${col_of[c]}$3:${col_of[c]}${last}"
    S.cell(i, 1, company[c]); S.cell(i, 2, c)
    S.cell(i, 3, f'=COUNT({col})')
    S.cell(i, 4, f'=COUNTIF({col},0)')
    S.cell(i, 5, f'=IFERROR(D{i}/C{i},0)')
    S.cell(i, 6, f'=IFERROR(COUNTIFS({col},0,{BUCKET},"Top 1K")/COUNTIFS({col},"<>",{BUCKET},"Top 1K"),0)')
    S.cell(i, 7, f'=IFERROR((COUNTIFS({col},0,{BUCKET},"Top 1K")+COUNTIFS({col},0,{BUCKET},"Top 5K")+COUNTIFS({col},0,{BUCKET},"Top 10K"))/(COUNTIFS({col},"<>",{BUCKET},"Top 1K")+COUNTIFS({col},"<>",{BUCKET},"Top 5K")+COUNTIFS({col},"<>",{BUCKET},"Top 10K")),0)')
    for j in range(1, 8):
        S.cell(i, j).font = body
    for j in (3, 4): S.cell(i, j).number_format = '#,##0'
    for j in (5, 6, 7): S.cell(i, j).number_format = '0.0%'
end1 = 4 + len(crawl_cols)
S.conditional_formatting.add(f'E5:G{end1}', CellIsRule(operator='greaterThanOrEqual', formula=['0.15'], fill=red, font=red_f))
for j, w in enumerate([14, 26, 13, 13, 14, 14, 14], 1):
    S.column_dimensions[L(j)].width = w

# Category x key crawler matrix
key = ['Applebot', 'Applebot-Extended', 'Googlebot', 'Google-Extended', 'GPTBot', 'OAI-SearchBot', 'Bingbot', 'PerplexityBot', 'ShapBot', 'ExaSearchBot', 'Brave (Googlebot rules)', 'YouBot', 'ClaudeBot', 'CCBot']
r0 = end1 + 3
S.cell(r0, 1, '% of sites blocking each crawler, by category').font = Font(name=F, size=14, bold=True)
S.cell(r0 + 1, 1, 'Category'); S.cell(r0 + 1, 2, 'Sites')
for k, c in enumerate(key):
    S.cell(r0 + 1, 3 + k, c)
for j in range(1, 3 + len(key)):
    cc = S.cell(r0 + 1, j); cc.font = hdr_font; cc.fill = hdr_fill; cc.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
S.row_dimensions[r0 + 1].height = 32
cats = d.category.value_counts().index.tolist()
for i, cat in enumerate(cats, r0 + 2):
    S.cell(i, 1, cat).font = body
    S.cell(i, 2, f'=COUNTIF({CAT},A{i})').number_format = '#,##0'
    for k, c in enumerate(key):
        col = f"Sites!${col_of[c]}$3:${col_of[c]}${last}"
        cell = S.cell(i, 3 + k, f'=IFERROR(COUNTIFS({CAT},$A{i},{col},0)/COUNTIFS({CAT},$A{i},{col},"<>"),0)')
        cell.number_format = '0.0%'; cell.font = body
end2 = r0 + 1 + len(cats)
S.conditional_formatting.add(f'C{r0+2}:{L(2+len(key))}{end2}', CellIsRule(operator='greaterThanOrEqual', formula=['0.2'], fill=red, font=red_f))
S.conditional_formatting.add(f'C{r0+2}:{L(2+len(key))}{end2}', CellIsRule(operator='lessThan', formula=['0.05'], fill=green, font=green_f))
for k in range(len(key)):
    S.column_dimensions[L(3 + k)].width = max(S.column_dimensions[L(3 + k)].width or 0, 12)
S.column_dimensions['A'].width = 30
S.freeze_panes = 'A5'

# ---------------- Crawlers reference ----------------
Cw = wb.create_sheet('Crawlers')
ref = [
    ('Applebot', 'Apple', 'Applebot', 'Search crawler for Siri, Spotlight and Safari suggestions.', 'If robots.txt does not name Applebot but names Googlebot, Applebot follows the Googlebot rules (per Apple).', 'https://support.apple.com/en-us/119829'),
    ('Applebot-Extended', 'Apple', 'Applebot-Extended', 'Control token: whether Apple may use content to train its generative AI models. Does not crawl by itself.', 'Falls back to the * group.', 'https://support.apple.com/en-us/119829'),
    ('Googlebot', 'Google', 'Googlebot', 'Google Search crawler.', '', 'https://developers.google.com/crawling/docs/crawlers-fetchers/google-common-crawlers'),
    ('Google-Extended', 'Google', 'Google-Extended', 'Control token for use of content in Gemini model training and grounding. Does not affect Google Search.', '', 'https://developers.google.com/crawling/docs/crawlers-fetchers/google-common-crawlers'),
    ('GPTBot', 'OpenAI', 'GPTBot', 'Collects content that may be used to train OpenAI models.', '', 'https://platform.openai.com/docs/bots'),
    ('OAI-SearchBot', 'OpenAI', 'OAI-SearchBot', 'Indexes sites for ChatGPT search results.', '', 'https://platform.openai.com/docs/bots'),
    ('ChatGPT-User', 'OpenAI', 'ChatGPT-User', 'Fetches pages when a ChatGPT user or GPT action requests them.', 'User-initiated agents are treated differently by some operators; the cell shows what the site\'s robots.txt says.', 'https://platform.openai.com/docs/bots'),
    ('Bingbot', 'Microsoft', 'bingbot (falls back to msnbot)', 'Bing search crawler; Bing\'s index also powers Copilot and many third-party search APIs.', 'Uses an msnbot group if no bingbot group exists.', 'https://www.bing.com/webmasters/help/which-crawlers-does-bing-use-8c184ec0'),
    ('PerplexityBot', 'Perplexity', 'PerplexityBot', 'Indexes sites for Perplexity search answers.', '', 'https://docs.perplexity.ai/guides/bots'),
    ('Perplexity-User', 'Perplexity', 'Perplexity-User', 'Fetches pages in response to a user\'s question.', 'See note on user-initiated agents.', 'https://docs.perplexity.ai/guides/bots'),
    ('ShapBot', 'Parallel', 'ShapBot', 'Parallel Web Systems crawler for its web search/research APIs.', '', 'https://docs.parallel.ai/resources/crawler'),
    ('ExaSearchBot', 'Exa', 'ExaSearchBot', 'Exa crawler for its AI search API.', '', 'https://crawler.exa.ai/'),
    ('Brave (Googlebot rules)', 'Brave', '(none: uses Googlebot rules)', 'Brave Search does not use its own user-agent token; it does not crawl pages Googlebot is not allowed to crawl.', 'Value is identical to Googlebot by construction.', 'https://search.brave.com/help/brave-search-crawler'),
    ('YouBot', 'You.com', 'YouBot', 'You.com crawler for its search and AI products.', '', 'https://you.com/docs/youbot'),
    ('ClaudeBot', 'Anthropic', 'ClaudeBot', 'Collects content that may be used for model training.', '', 'https://support.claude.com'),
    ('Claude-SearchBot', 'Anthropic', 'Claude-SearchBot', 'Indexes content to improve Claude\'s search results.', '', 'https://support.claude.com'),
    ('Claude-User', 'Anthropic', 'Claude-User', 'Fetches pages when a Claude user asks.', 'See note on user-initiated agents.', 'https://support.claude.com'),
    ('CCBot', 'Common Crawl', 'CCBot', 'Builds the open Common Crawl corpus, widely used for AI training.', '', 'https://commoncrawl.org/ccbot'),
    ('Meta-ExternalAgent', 'Meta', 'meta-externalagent', 'Meta crawler for AI training and products.', '', 'https://developers.facebook.com/docs/sharing/webmasters/web-crawlers'),
    ('Amazonbot', 'Amazon', 'Amazonbot', 'Amazon crawler (Alexa, Rufus and other services).', '', 'https://developer.amazon.com/amazonbot'),
    ('Bytespider', 'ByteDance', 'Bytespider', 'ByteDance crawler (TikTok/Doubao AI).', '', ''),
    ('DuckDuckBot', 'DuckDuckGo', 'DuckDuckBot', 'DuckDuckGo search crawler.', '', 'https://duckduckgo.com/duckduckgo-help-pages/results/duckduckbot'),
    ('DuckAssistBot', 'DuckDuckGo', 'DuckAssistBot', 'Fetches pages for DuckDuckGo\'s AI-assisted answers.', '', 'https://duckduckgo.com/duckduckgo-help-pages/results/duckassistbot'),
    ('YandexBot', 'Yandex', 'YandexBot (falls back to Yandex)', 'Yandex search crawler.', 'Uses a "Yandex" group if no YandexBot group exists.', ''),
    ('Baiduspider', 'Baidu', 'Baiduspider', 'Baidu search crawler.', '', ''),
    ('Yeti', 'Naver', 'Yeti', 'Naver search crawler (Korea).', '', ''),
    ('PetalBot', 'Huawei', 'PetalBot', 'Huawei Petal Search crawler.', '', ''),
    ('cohere-ai', 'Cohere', 'cohere-ai', 'Cohere AI agent token.', '', ''),
    ('MistralAI-User', 'Mistral', 'MistralAI-User', 'Fetches pages for Mistral Le Chat users.', '', ''),
    ('Diffbot', 'Diffbot', 'Diffbot', 'Knowledge-graph / structured web data crawler.', '', ''),
    ('All bots (*)', 'Default rule', '*', 'The default group that applies to any crawler not named in robots.txt.', '', 'https://www.rfc-editor.org/rfc/rfc9309'),
]
hd = ['Column', 'Company', 'robots.txt token evaluated', 'What it does', 'Fallback / notes', 'Documentation']
for j, h in enumerate(hd, 1):
    c = Cw.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill
for i, row in enumerate(ref, 2):
    for j, v in enumerate(row, 1):
        c = Cw.cell(i, j, v); c.font = body; c.alignment = Alignment(wrap_text=True, vertical='top')
for j, w in enumerate([22, 14, 26, 50, 42, 48], 1):
    Cw.column_dimensions[L(j)].width = w
Cw.freeze_panes = 'A2'

# ---------------- Categories ----------------
Cg = wb.create_sheet('Categories')
for j, h in enumerate(['Category', 'Sites', '% of 50K', 'Manually reviewed'], 1):
    c = Cg.cell(1, j, h); c.font = hdr_font; c.fill = hdr_fill
SRC = f"Sites!$F$3:$F${last}"
for i, cat in enumerate(cats, 2):
    Cg.cell(i, 1, cat).font = body
    Cg.cell(i, 2, f'=COUNTIF({CAT},A{i})').number_format = '#,##0'
    Cg.cell(i, 3, f'=B{i}/{N}').number_format = '0.0%'
    Cg.cell(i, 4, f'=COUNTIFS({CAT},A{i},{SRC},"Manual review")').number_format = '#,##0'
Cg.column_dimensions['A'].width = 34
for col in 'BCD': Cg.column_dimensions[col].width = 16

# ---------------- Method & notes ----------------
M = wb.create_sheet('Method & notes')
notes = [
    ('What this is', 'robots.txt rules for the top 50,000 website origins, evaluated for 30 crawlers plus the default (*) rule. Collected 5 Oct 2026.'),
    ('Site list', 'Top 50,000 origins from the Chrome UX Report (CrUX) global list for Aug 2026 (via github.com/zakird/crux-top-lists). CrUX ranks origins (e.g. m.youtube.com and www.youtube.com separately) by real Chrome page loads, in buckets (Top 1K, 5K, 10K, 50K). robots.txt is per origin, so origin is the right unit.'),
    ('Traffic measure', 'No free source publishes yearly visit counts for 50K sites (Similarweb-style visit data is paid). Popularity is shown as: (1) the CrUX bucket from Chrome traffic, (2) the Tranco rank of the domain (research ranking aggregating several lists over 30 days; list ID 647LX, 5 Oct 2026), and (3) "Popularity rank (est.)", which sorts by CrUX bucket and then Tranco rank. Note: Tranco breaks ties alphabetically beyond roughly rank 50,000, so treat Tranco ranks there as approximate.'),
    ('Allowed / Blocked', '1 (Allowed) = the robots.txt rule that applies to this crawler permits fetching both the homepage "/" and an ordinary content URL. 0 (Blocked) = either is disallowed, i.e. a site-wide or near site-wide block. Path-level restrictions (e.g. /search, /account) do not count as blocked. A blank = no robots.txt data for that site.'),
    ('Parsing rules', 'RFC 9309 with Google matching semantics: the most specific user-agent group applies (otherwise the * group); within a group the longest matching Allow/Disallow wins, ties go to Allow; * and $ wildcards supported; groups with the same user-agent are merged. A missing robots.txt (HTTP 404/410 or an HTML page served instead) means everything is allowed.'),
    ('Crawler-specific fallbacks', 'Applebot uses the Googlebot group when Applebot is not named (Apple\'s documented behavior). Brave has no token of its own and follows Googlebot rules, so its column equals Googlebot. Bingbot falls back to an msnbot group; YandexBot falls back to a Yandex group.'),
    ('How robots.txt was fetched', 'Fetched live on 5 Oct 2026 from GitHub-hosted runners (39,852 sites). When the live fetch was blocked or failed (bot walls, 403s, timeouts), the most recent Common Crawl capture (Jul–Sep 2026 crawls) was used instead (2,982 sites). The "robots.txt source" column records which one applies to each row.'),
    ('No data', '7,166 sites (14%) have no robots.txt data: they block automated requests at the firewall (Cloudflare, Akamai, DDoS-Guard and similar), were unreachable or geo-blocked, and had no usable Common Crawl capture. Their crawler cells are blank, and the Summary percentages exclude them.'),
    ('Important caveat', 'robots.txt shows what a site asks crawlers to do; it is not the whole picture. Many sites also block bots at the network level (e.g. Cloudflare\'s AI-crawler blocking) regardless of robots.txt, and some sites serve different robots.txt files to different requesters or have separate agreements (e.g. Reddit\'s public robots.txt disallows all crawlers, while it licenses access to specific partners). How operators treat user-initiated fetchers (ChatGPT-User, Perplexity-User, Claude-User) also varies.'),
    ('Categories', 'The top 10,000 sites were categorized by manual review of the domain, page title and description. The remaining 40,000 were categorized by a classifier trained on those 10,000 labels (multilingual text embeddings of title/description plus domain and keyword features); its accuracy in cross-validation on the reviewed sites was about 79%. Domain rules (.gov/.gob/.gouv, .edu/.ac) and the UT1 web category blocklists (Université Toulouse Capitole) override the model where they apply. The "Category source" column shows which method applies, with a confidence band for model predictions. Categories beyond your list were added where they emerged: Gambling & Lottery, Anime/Manga & Comics, Entertainment, Business & Professional Services, Productivity & Workplace Tools, File Sharing & Downloads, AI Tools & Chatbots, Dating, Religion, Weather, Books & Literature, Lifestyle & Fashion, Science, Utilities & Telecom, and Unclassified / Parked.'),
    ('Country', 'If the site uses a country-code domain (e.g. .jp, .co.uk), that country. Otherwise, the country where the site draws the most Chrome traffic, estimated from CrUX per-country top lists (rank bucket in each of 238 countries, weighted by each country\'s Chrome user base). Sites in the top 10K of 40+ countries with no country-code domain are labeled "Global"; their largest markets are in "Top markets (est.)". Generic ccTLDs (.io, .co, .tv, .me, .ai, etc.) are not treated as countries.'),
    ('Reproducibility', 'Crawl code, raw robots.txt bodies and Common Crawl fallback data: github.com/jomilu93/robots-txt-crawl.'),
]
M['A1'] = 'Method & notes'; M['A1'].font = Font(name=F, size=14, bold=True)
for i, (k, v) in enumerate(notes, 3):
    M.cell(i, 1, k).font = bold; M.cell(i, 1).alignment = Alignment(vertical='top')
    M.cell(i, 2, v).font = body; M.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='top')
M.column_dimensions['A'].width = 24; M.column_dimensions['B'].width = 120

wb.move_sheet('Summary', offset=-1)
wb.active = 0
wb.save(OUT)
print('saved', OUT)
